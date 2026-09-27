"""Check literal workbook conversion and conservative subject enrichment."""

import asyncio
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import httpx
from openpyxl import load_workbook

from app.config import Settings
from app.ingestion.visual_workbook import (
    TRANSCRIPTION_PROMPT, VisualCell, VisualSheet, VisualWorkbook, build_workbook,
)
from app.ingestion.workbook_evidence import keep_visible_values
from app.llm.extractor import ExtractionError
from app.main import app
from app.resolvers.subject_embeddings import embedding_subject_matches
from app.schemas.context import RoutineContext, SubjectRecord
from app.schemas.extraction import RoutineExtraction
from app.services.routine_processor import process_routine


class FakeEmbeddings:
    def __init__(self, vectors):
        self.vectors = vectors
        self.calls = []

    async def aembed_documents(self, texts):
        self.calls.append(texts)
        return [self.vectors[text] for text in texts]


class VisualWorkbookTest(unittest.TestCase):
    def test_image_upload_saves_literal_excel_then_returns_grounded_json(self):
        class FakeModel:
            def __init__(self):
                self.responses = [
                    {"sheets": [{"cells": [
                        {"row": 1, "column": 1, "text": "Dept of AEIE"},
                        {"row": 2, "column": 1, "text": "Section: AEIE.1"},
                        {"row": 3, "column": 1, "text": "MON"},
                        {"row": 3, "column": 2, "text": "Network Analysis MP", "column_span": 2},
                    ]}]},
                    {"routines": [{"department": "AEIE", "section": "AEIE.1",
                                   "default_room": "R-999", "slots": [{
                                       "day": "Monday", "slot_type": "class",
                                       "activities": [{"subject_raw": "Network Analysis",
                                                       "faculty_raw": ["MP", "NF9"]}],
                                   }]}]},
                ]
                self.calls = []

            async def ainvoke(self, messages):
                self.calls.append(messages)
                return SimpleNamespace(content=json.dumps(self.responses.pop(0)))

        class EmptyMaster:
            async def load(self, routine, college_id):
                return RoutineContext(subjects=[], faculty=[], groups=[], sections=[], college_id=college_id)

        model = FakeModel()
        with tempfile.TemporaryDirectory() as directory:
            app.state.master_api = EmptyMaster()
            with patch("app.main.Settings", return_value=Settings(output_dir=directory, openai_api_key="test")), patch(
                "app.services.routine_processor.image_data_urls", return_value=["data:image/png;base64,AA=="],
            ), patch("app.llm.extractor.ChatOpenAI", return_value=model):
                async def request():
                    transport = httpx.ASGITransport(app=app)
                    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                        return await client.post("/routines/standardize?college_id=1", files={
                            "file": ("aeie.png", b"fake-image", "image/png"),
                        })
                response = asyncio.run(request())
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(len(model.calls), 2)
            self.assertEqual(model.calls[0][0]["content"][0]["type"], "text")
            self.assertEqual(model.calls[1][0]["content"][1]["type"], "text")
            self.assertEqual(payload["routines"][0]["department"], "AEIE")
            self.assertIsNone(payload["routines"][0]["default_room"])
            faculty = payload["routines"][0]["slots"][0]["activities"][0]["faculty"]
            self.assertEqual([item["raw"] for item in faculty], ["MP"])
            self.assertTrue(payload["conversion_review_reasons"])
            self.assertEqual(json.loads(Path(payload["output_file"]).read_text()), payload)
            workbook = load_workbook(payload["converted_workbook_file"])
            try:
                self.assertEqual(workbook.active["C3"].value, None)
                self.assertEqual(workbook.active["B3"].value, "Network Analysis MP")
                self.assertNotIn("R-999", [cell.value for row in workbook.active for cell in row])
            finally:
                workbook.close()

    def test_transcribed_cells_and_merges_survive_workbook_round_trip(self):
        self.assertIn("never guess", TRANSCRIPTION_PROMPT)
        transcription = VisualWorkbook(sheets=[VisualSheet(cells=[
            VisualCell(row=1, column=1, text="Section: AIML.2A", column_span=3),
            VisualCell(row=2, column=1, text="MON"),
            VisualCell(row=2, column=2, text="=AIML_SC\nDSA Lab", column_span=2),
        ])])
        data = build_workbook(transcription, Settings())
        workbook = load_workbook(BytesIO(data))
        try:
            sheet = workbook.active
            self.assertEqual(sheet["B2"].value, "=AIML_SC\nDSA Lab")
            self.assertEqual(sheet["B2"].data_type, "s")
            self.assertEqual({str(area) for area in sheet.merged_cells.ranges}, {"A1:C1", "B2:C2"})
        finally:
            workbook.close()

    def test_visible_page_heading_and_period_header_survive_conversion(self):
        self.assertIn("ENTIRE page", TRANSCRIPTION_PROMPT)
        self.assertIn("period number or name", TRANSCRIPTION_PROMPT)
        transcription = VisualWorkbook(sheets=[VisualSheet(cells=[
            VisualCell(row=1, column=1, text="Visible college heading", column_span=3),
            VisualCell(row=2, column=1, text="Section A"),
            VisualCell(row=3, column=2, text="1\n9:20 - 10:10"),
            VisualCell(row=4, column=1, text="Mo"),
            VisualCell(row=4, column=2, text="DSA"),
        ])])
        workbook = load_workbook(BytesIO(build_workbook(transcription, Settings())))
        try:
            sheet = workbook.active
            self.assertEqual(sheet["A1"].value, "Visible college heading")
            self.assertEqual(sheet["B3"].value, "1\n9:20 - 10:10")
            self.assertEqual(sheet["B4"].value, "DSA")
        finally:
            workbook.close()

    def test_overlapping_transcribed_cells_are_rejected(self):
        transcription = VisualWorkbook(sheets=[VisualSheet(cells=[
            VisualCell(row=1, column=1, text="A", column_span=2),
            VisualCell(row=1, column=2, text="B"),
        ])])
        with self.assertRaises(ExtractionError):
            build_workbook(transcription, Settings())

    def test_unseen_raw_values_become_null_before_enrichment(self):
        routine = RoutineExtraction.model_validate({
            "department": "Imaginary", "section": "AIML.2A", "slots": [{
                "day": "Monday", "slot_type": "class", "activities": [{
                    "subject_raw": "DSA Lab", "faculty_raw": ["AIML_SC", "NF9"],
                    "room_raw": "R-999",
                }],
            }],
        })
        corrected, reasons = keep_visible_values(
            [routine], ["Section: AIML.2A", "MON", "DSA Lab AIML_SC"],
        )
        self.assertIsNone(corrected[0].department)
        activity = corrected[0].slots[0].activities[0]
        self.assertEqual(activity.subject_raw, "DSA Lab")
        self.assertEqual(activity.faculty_raw, ["AIML_SC"])
        self.assertIsNone(activity.room_raw)
        self.assertEqual(len(reasons), 3)

    def test_unseen_day_period_and_time_are_cleared(self):
        routine = RoutineExtraction.model_validate({"slots": [{
            "day": "Friday", "start_period": 4, "end_period": 5,
            "start_time": "11:50", "end_time": "13:30",
            "slot_type": "class", "activities": [{"subject_raw": "DSA"}],
        }]})
        corrected, reasons = keep_visible_values([routine], ["MON", "1 9.20am-10.10am", "DSA"])
        slot = corrected[0].slots[0]
        self.assertIsNone(slot.day)
        self.assertIsNone(slot.start_period)
        self.assertIsNone(slot.end_period)
        self.assertIsNone(slot.start_time)
        self.assertIsNone(slot.end_time)
        self.assertEqual(len(reasons), 5)

    def test_two_letter_day_labels_in_converted_workbook_are_accepted(self):
        days = [("Monday", "Mo"), ("Tuesday", "Tu"), ("Wednesday", "We"),
                ("Thursday", "Th"), ("Friday", "Fr")]
        routine = RoutineExtraction.model_validate({"slots": [
            {"day": day, "activities": []} for day, _ in days
        ]})
        corrected, reasons = keep_visible_values([routine], [label for _, label in days])
        self.assertEqual([slot.day for slot in corrected[0].slots], [day for day, _ in days])
        self.assertEqual(reasons, [])

    def test_one_routines_values_cannot_be_justified_by_another_sheet(self):
        routines = [RoutineExtraction.model_validate({
            "department": "CSE", "slots": [{"activities": [{"subject_raw": "DSA"}]}],
        }), RoutineExtraction.model_validate({
            "department": "ECE", "slots": [{"activities": [{"subject_raw": "DSA"}]}],
        })]
        corrected, reasons = keep_visible_values(routines, [
            ["Department: CSE", "DSA"], ["Department: ECE", "Circuits"],
        ])
        self.assertEqual(corrected[0].slots[0].activities[0].subject_raw, "DSA")
        self.assertIsNone(corrected[1].slots[0].activities[0].subject_raw)
        self.assertEqual(len(reasons), 1)

    def test_embedding_fallback_batches_and_requires_clear_winner(self):
        routines = [RoutineExtraction.model_validate({
            "department": "CSE", "slots": [{"activities": [
                {"subject_raw": "DSA"}, {"subject_raw": "DSA"},
            ]}],
        })]
        context = RoutineContext(subjects=[
            SubjectRecord(id=1, name="Data Structures and Algorithms", stream="CSE"),
            SubjectRecord(id=2, name="Digital Systems", stream="CSE"),
            SubjectRecord(id=3, name="DSA", stream="ECE"),
        ], faculty=[], groups=[], sections=[])
        model = FakeEmbeddings({
            "Data Structures and Algorithms": [1.0, 0.0],
            "Digital Systems": [0.0, 1.0], "DSA": [0.99, 0.01],
        })
        result = asyncio.run(embedding_subject_matches(
            routines, context, Settings(subject_embedding_model="test"), model,
        ))
        self.assertEqual(len(model.calls), 1)
        self.assertEqual(model.calls[0], ["Data Structures and Algorithms", "Digital Systems", "DSA"])
        self.assertEqual(next(iter(result.values())).subject_master_id, 1)
        self.assertEqual(next(iter(result.values())).match_method, "embedding")

        close_model = FakeEmbeddings({
            "Data Structures and Algorithms": [1.0, 0.0],
            "Digital Systems": [0.99, 0.01], "DSA": [1.0, 0.0],
        })
        ambiguous = asyncio.run(embedding_subject_matches(
            routines, context, Settings(subject_embedding_model="test"), close_model,
        ))
        self.assertEqual(ambiguous, {})

    def test_shared_processor_marks_embedding_subject_id_for_review(self):
        routine = RoutineExtraction.model_validate({
            "department": "CSE", "slots": [{
                "day": "Monday", "slot_type": "class",
                "activities": [{"subject_raw": "DSA"}],
            }],
        })

        class Master:
            async def load(self, extracted, college_id):
                return RoutineContext(subjects=[
                    SubjectRecord(id=8, name="Data Structures and Algorithms", stream="CSE"),
                ], faculty=[], groups=[], sections=[], college_id=college_id)

        model = FakeEmbeddings({"Data Structures and Algorithms": [1.0, 0.0], "DSA": [1.0, 0.0]})
        settings = Settings(subject_embedding_model="test", subject_embedding_api_key="test")
        with patch("app.services.routine_processor.is_excel_workbook", return_value=True), patch(
            "app.services.routine_processor.TINTProfile.extract", return_value=[routine],
        ), patch("app.resolvers.subject_embeddings.OpenAIEmbeddings", return_value=model):
            result = asyncio.run(process_routine(b"workbook", settings, Master(), 2))
        activity = result.routines[0].slots[0].activities[0]
        self.assertEqual(activity.subject.subject_master_id, 8)
        self.assertEqual(activity.subject.match_method, "embedding")
        self.assertTrue(activity.requires_review)
        self.assertIn("Subject matched semantically", activity.review_reasons[0])


if __name__ == "__main__":
    unittest.main()
