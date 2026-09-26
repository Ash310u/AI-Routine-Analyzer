"""Turn Excel timetable cells into one location-aware model input."""

from datetime import date, datetime, time
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from app.config import Settings
from app.ingestion.document import DocumentError


OLE_SIGNATURE = bytes.fromhex("D0 CF 11 E0 A1 B1 1A E1")


def is_excel_workbook(data: bytes) -> bool:
    if data.startswith(OLE_SIGNATURE):
        return True
    if not data.startswith(b"PK"):
        return False
    try:
        with ZipFile(BytesIO(data)) as archive:
            return "xl/workbook.xml" in archive.namelist()
    except BadZipFile:
        return False


def workbook_text(data: bytes, settings: Settings) -> str:
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise DocumentError(f"File exceeds {settings.max_upload_mb} MB limit")
    if data.startswith(OLE_SIGNATURE):
        return _xls_text(data, settings)
    if data.startswith(b"PK"):
        return _xlsx_text(data, settings)
    raise DocumentError("Not an Excel workbook")


def _xlsx_text(data: bytes, settings: Settings) -> str:
    from openpyxl import load_workbook

    try:
        with ZipFile(BytesIO(data)) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise DocumentError("Upload an Excel workbook, not another ZIP file")
            expanded = sum(entry.file_size for entry in archive.infolist())
            if expanded > 50 * 1024 * 1024:
                raise DocumentError("Expanded workbook is too large")
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise DocumentError("Invalid Excel workbook") from exc
    try:
        parts = []
        count = 0
        for sheet in workbook:
            if sheet.sheet_state != "visible":
                continue
            parts.append(f"Sheet: {sheet.title}")
            if sheet.merged_cells.ranges:
                parts.append("Merged cells: " + ", ".join(str(area) for area in sheet.merged_cells.ranges))
            for row in sheet:
                for cell in row:
                    if cell.value is None or sheet.row_dimensions[cell.row].hidden or sheet.column_dimensions[cell.column_letter].hidden:
                        continue
                    count += 1
                    if count > settings.max_workbook_cells:
                        raise DocumentError("Workbook has too many populated cells")
                    parts.append(f"{cell.coordinate}: {_value(cell.value)}")
        return _finish(parts, settings)
    finally:
        workbook.close()


def _xls_text(data: bytes, settings: Settings) -> str:
    import xlrd

    try:
        workbook = xlrd.open_workbook(file_contents=data, formatting_info=True)
    except (xlrd.XLRDError, OSError, ValueError) as exc:
        raise DocumentError("Invalid legacy Excel workbook") from exc
    parts = []
    count = 0
    for sheet in workbook.sheets():
        if sheet.visibility != 0:
            continue
        parts.append(f"Sheet: {sheet.name}")
        if sheet.merged_cells:
            parts.append("Merged cells: " + ", ".join(
                f"{_column(c0)}{r0 + 1}:{_column(c1 - 1)}{r1}"
                for r0, r1, c0, c1 in sheet.merged_cells
            ))
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                cell = sheet.cell(row, col)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                count += 1
                if count > settings.max_workbook_cells:
                    raise DocumentError("Workbook has too many populated cells")
                value = xlrd.xldate_as_datetime(cell.value, workbook.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                parts.append(f"{_column(col)}{row + 1}: {_value(value)}")
    return _finish(parts, settings)


def _column(index: int) -> str:
    label = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        label = chr(65 + remainder) + label
    return label


def _value(value: object) -> str:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).replace("\n", " / ").strip()


def _finish(parts: list[str], settings: Settings) -> str:
    if not parts:
        raise DocumentError("Workbook has no visible timetable data")
    text = "\n".join(parts)
    if len(text) > settings.max_workbook_characters:
        raise DocumentError("Workbook text is too large for one extraction request")
    return text
