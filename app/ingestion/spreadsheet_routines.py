"""Extract timetable blocks from structured Excel grids without an LLM call."""

import re
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO

from app.config import Settings
from app.ingestion.document import DocumentError
from app.ingestion.spreadsheet import OLE_SIGNATURE, workbook_text
from app.schemas.extraction import ActivityExtraction, RoutineExtraction, SlotExtraction


DAY_NAMES = {
    "MON": "Monday", "TUE": "Tuesday", "WED": "Wednesday",
    "THU": "Thursday", "FRI": "Friday", "SAT": "Saturday", "SUN": "Sunday",
}
CLOCK = re.compile(r"(\d{1,2}):(\d{2})\s*(am|pm)?", re.I)
PERIOD = re.compile(r"\b(?:period|p)\s*[-:#]?\s*(\d{1,2})\b", re.I)
GROUP = re.compile(r"\b(?:gr|group)[\s-]*([a-z])\b", re.I)
CODE = re.compile(r"\b[A-Z]{2,}(?:[- ]?[A-Z]{1,5})?[- ]?\d{3}[A-Z]?\b")
ROOM = re.compile(r"\b(?:R[\s-]?\d{2,4}|Room\s*[:#]?\s*[A-Z]?\d+)\b", re.I)
INITIALS = re.compile(r"[\[(]([A-Z0-9+/&,\s-]+)[\])]")
VENUE = re.compile(r"[\[(]((?:H/W\s+Lab|Project\s+Lab[-\s]?\d*|PL[-\s]?\d+|Central\s+Computing\s+Lab))[\])]", re.I)


@dataclass
class Grid:
    title: str
    rows: int
    columns: int
    values: dict[tuple[int, int], object]
    merges: list[tuple[int, int, int, int]]  # min row, max row, min col, max col

    def value(self, row: int, column: int) -> object | None:
        return self.values.get((row, column))


class NoRoutineBlocks(DocumentError):
    """The workbook uses a layout other than the supported block grid."""


def extract_workbook_routines(data: bytes, settings: Settings) -> list[RoutineExtraction]:
    # This also checks the upload, expansion, cell, and text size limits.
    workbook_text(data, settings)
    grids = _legacy_grids(data) if data.startswith(OLE_SIGNATURE) else _modern_grids(data)
    routines = []
    for grid in grids:
        headers = [row for row in range(1, grid.rows + 1)
                   if "department:" in str(grid.value(row, 1) or "").casefold()]
        for index, row in enumerate(headers):
            end = headers[index + 1] if index + 1 < len(headers) else grid.rows + 1
            routines.append(_extract_block(grid, row, end))
    if not routines:
        raise NoRoutineBlocks("Workbook contains no timetable blocks with a Department header")
    return routines


def _modern_grids(data: bytes) -> list[Grid]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    try:
        grids = []
        for sheet in workbook:
            if sheet.sheet_state != "visible":
                continue
            values = {(cell.row, cell.column): cell.value for row in sheet for cell in row
                      if cell.value is not None
                      and not sheet.row_dimensions[cell.row].hidden
                      and not sheet.column_dimensions[cell.column_letter].hidden}
            merges = [(area.min_row, area.max_row, area.min_col, area.max_col)
                      for area in sheet.merged_cells.ranges]
            grids.append(Grid(sheet.title, sheet.max_row, sheet.max_column, values, merges))
        return grids
    finally:
        workbook.close()


def _legacy_grids(data: bytes) -> list[Grid]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data, formatting_info=True)
    grids = []
    for sheet in workbook.sheets():
        if sheet.visibility != 0:
            continue
        values = {}
        for row in range(sheet.nrows):
            for column in range(sheet.ncols):
                cell = sheet.cell(row, column)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                value = (xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                         if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
                values[(row + 1, column + 1)] = value
        merges = [(r0 + 1, r1, c0 + 1, c1) for r0, r1, c0, c1 in sheet.merged_cells]
        grids.append(Grid(sheet.name, sheet.nrows, sheet.ncols, values, merges))
    return grids


def _extract_block(grid: Grid, header_row: int, end_row: int) -> RoutineExtraction:
    header = str(grid.value(header_row, 1) or "")
    time_row = None
    ranges = {}
    periods = {}
    for row in range(header_row + 1, min(header_row + 4, end_row)):
        candidate = {}
        candidate_periods = {}
        previous_end = None
        for column in range(2, grid.columns + 1):
            cell = grid.value(row, column)
            if cell is None:
                continue
            parsed = _time_range(str(cell), previous_end)
            if parsed:
                candidate[column] = parsed
                period = PERIOD.search(str(cell))
                if period:
                    candidate_periods[column] = int(period.group(1))
                previous_end = parsed[1]
        if len(candidate) >= 3:
            time_row, ranges = row, candidate
            periods = candidate_periods
            break
    if time_row is None:
        raise DocumentError(f"No time header found in {grid.title} near row {header_row}")

    merged = {}
    for first_row, last_row, first_col, last_col in grid.merges:
        if last_row < time_row + 1 or first_row >= end_row:
            continue
        for row in range(max(first_row, time_row + 1), min(last_row + 1, end_row)):
            for column in range(first_col, last_col + 1):
                merged[(row, column)] = (first_row, first_col, last_col)

    slots = []
    for row in range(time_row + 1, end_row):
        day = DAY_NAMES.get(str(grid.value(row, 1) or "").strip().upper())
        if day is None:
            continue
        for column, (start, default_end) in ranges.items():
            anchor_row, anchor_col, last_col = merged.get((row, column), (row, column, column))
            if anchor_col != column:
                continue
            value = grid.value(anchor_row, anchor_col)
            if value is None or not str(value).strip():
                continue
            end = ranges.get(last_col, (start, default_end))[1]
            start_period = periods.get(column)
            end_period = periods.get(last_col)
            raw = str(value).strip()
            if raw.casefold() == "break":
                slot = SlotExtraction(day=day, start_period=start_period, end_period=end_period,
                                      start_time=start, end_time=end,
                                      slot_type="break", activities=[])
            else:
                chunks = [part.strip() for part in re.split(r"\n| {3,}", raw) if part.strip()]
                if len(GROUP.findall(raw)) < 2 or not all(GROUP.search(part) for part in chunks):
                    chunks = [raw]
                activities = [_activity(part, grid.title, anchor_row, anchor_col) for part in chunks]
                slot = SlotExtraction(day=day, start_period=start_period, end_period=end_period,
                                      start_time=start, end_time=end,
                                      slot_type="other" if "research hours" in raw.casefold() else "class",
                                      activities=activities)
            slots.append(slot)

    if not slots:
        raise DocumentError(f"No timetable slots found in {grid.title} near row {header_row}")
    title = str(grid.value(header_row - 1, 1) or "").splitlines()[0].strip() or None
    return RoutineExtraction(
        college=title,
        course="M.Tech" if "m.tech" in grid.title.casefold() else None,
        department=_metadata(header, "Department", ("Year", "Semester", "Section", "Room")),
        year=_metadata(header, "Year", ("Semester", "Section", "Room")),
        semester=_metadata(header, "Semester", ("Section", "Room")),
        section=_metadata(header, "Section", ("Room",)),
        default_room=_metadata(header, "Room", ()),
        routine_version=None,
        slots=slots,
    )


def _metadata(header: str, label: str, following: tuple[str, ...]) -> str | None:
    match = re.search(rf"\b{label}\s*[:\-]\s*", header, re.I)
    if not match:
        return None
    rest = header[match.end():]
    for next_label in following:
        next_match = re.search(rf"\b{next_label}\s*[:\-]", rest, re.I)
        if next_match:
            rest = rest[:next_match.start()]
    return rest.strip() or None


def _time_range(value: str, previous_end: time | None) -> tuple[time, time] | None:
    matches = list(CLOCK.finditer(value))
    if len(matches) < 2:
        return None
    start = _clock_value(matches[0], previous_end)
    end = _clock_value(matches[1], start, strictly_after=True)
    return (start, end) if start and end and start < end else None


def _clock_value(match: re.Match, after: time | None, strictly_after: bool = False) -> time | None:
    hour, minute = int(match.group(1)), int(match.group(2))
    if minute > 59 or hour > 23:
        return None
    hours = [hour] if hour > 12 else [hour % 12, hour % 12 + 12]
    candidates = [time(h, minute) for h in hours if 6 <= h <= 22]
    if after is not None:
        if strictly_after:
            candidates = [candidate for candidate in candidates if candidate > after]
        else:
            candidates = [candidate for candidate in candidates if candidate >= after]
    if not candidates:
        return None
    # A repeated, contradictory AM/PM label (for example 11:20pm after 10:25am)
    # must not turn a daytime class into an overnight timetable.
    suffix = (match.group(3) or "").casefold()
    return min(candidates, key=lambda candidate: (
        0 if not suffix or (candidate.hour < 12) == (suffix == "am") else 1,
        candidate.hour * 60 + candidate.minute - (
            after.hour * 60 + after.minute if after else 8 * 60),
    ))


def _activity(raw: str, sheet: str, row: int, column: int) -> ActivityExtraction:
    groups = GROUP.findall(raw)
    codes = CODE.findall(raw.upper())
    faculty = []
    subject = raw
    for match in INITIALS.finditer(raw.upper()):
        tokens = re.split(r"[+/&,\s]+", match.group(1).strip())
        if tokens and all(re.fullmatch(r"[A-Z]{2,5}\d?", token) for token in tokens) and not any(
            token in {"LAB", "ROOM", "CLASS", "SEC"} for token in tokens
        ):
            faculty.extend(tokens)
            subject = subject.replace(raw[match.start():match.end()], "", 1)
    faculty = list(dict.fromkeys(faculty))
    room = ROOM.search(raw)
    venue = VENUE.search(raw)
    if room:
        subject = ROOM.sub("", subject, count=1)
    if venue:
        subject = subject.replace(venue.group(0), "", 1)
    subject = GROUP.sub("", subject)
    subject = re.sub(r"[\[(]\s*[\])]", "", subject)
    subject = re.sub(r"\s+", " ", subject).strip(" -()/") or None
    column_name = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        column_name = chr(65 + remainder) + column_name
    return ActivityExtraction(
        group_raw=f"Gr-{groups[0].upper()}" if len(set(g.upper() for g in groups)) == 1 else None,
        subject_raw=subject,
        subject_code_raw=codes[0].replace(" ", "") if codes else None,
        subject_type_raw="Lab" if "lab" in raw.casefold() else None,
        faculty_raw=faculty,
        room_raw=room.group(0) if room else venue.group(1) if venue else None,
        notes=f"Source: {sheet}!{column_name}{row}",
    )
