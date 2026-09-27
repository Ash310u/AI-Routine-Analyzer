"""Transcribe a visual timetable into literal Excel cells before routine parsing."""

import json
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from app.config import Settings
from app.llm.extractor import ExtractionError, ModelServiceError


TRANSCRIPTION_PROMPT = """Transcribe the supplied timetable PDF or image into Excel cells.

Your task is literal transcription only. Do not extract a routine schema here.
Do not explain, normalize, translate, expand abbreviations, infer headings,
calculate times, assign period numbers, supply missing data, or generate any
faculty, subject, section, or database IDs. Copy only text visibly present.
Keep punctuation, line breaks, initials, and numbers exactly as shown. If a
cell is blank or unreadable, omit it; never guess its contents. Preserve the
visual row/column order and only mark a span when the source visibly shows a
merged cell. Do not turn one merged cell into repeated values. Separate
visually independent timetable tables into separate sheets. Do not add sheet
titles, column labels, or Department headings that are absent in the source.
Examine the ENTIRE page for each table, not just the class/activity grid.
Read sideways or rotated pages in their upright reading orientation. Include
all legible timetable-related text around the grid in that table's sheet:
the institution/college heading, timetable title, printed or handwritten
version/date, course/department/year/semester/section/room labels, and any
visible table captions or explanatory notes. Put headings above the grid in
preceding worksheet rows, in their visual order; keep side headings beside
the grid. Transcribe the COMPLETE header row of the grid, including a day
heading if present and EVERY visible period number or name and its printed
start/end time, in the same columns as the corresponding activities. Keep
numbers and times as the source prints them; do not calculate missing ones.
Only copy a heading into multiple sheets when the source visibly repeats it
or clearly shows it applying to each table. If a heading, title, or period
label is absent or unreadable, omit that cell instead of creating it.
Ignore signatures, stamps, and approval marks outside the timetable.
Preserve handwritten timetable text when readable; do not silently replace
or complete an unclear printed entry.

Examples of forbidden changes: do not turn "2nd yr" into "B.Tech"; do not
expand "AB" into a teacher name; do not invent period 5 between visible
periods 4 and 6; do not copy text into an empty merged cell; do not turn a
faculty code into an employee ID. A later step will interpret the workbook.

Return only JSON in this shape:
{
  "sheets": [
    {
      "cells": [
        {"row": 1, "column": 1, "text": "exact visible text",
         "row_span": 1, "column_span": 1}
      ]
    }
  ]
}

Rows and columns are 1-based positions within each worksheet, including its
headings and period/time header row. A merged
cell appears once at its top-left coordinate with its visible row_span and
column_span. Include headings in cells only when they are visibly present.
If a page contains no timetable table, do not create a sheet for it.
"""


class VisualCell(BaseModel):
    model_config = ConfigDict(extra="forbid")

    row: int = Field(ge=1, le=1000)
    column: int = Field(ge=1, le=100)
    text: str = Field(min_length=1, max_length=32767)
    row_span: int = Field(default=1, ge=1, le=100)
    column_span: int = Field(default=1, ge=1, le=100)


class VisualSheet(BaseModel):
    model_config = ConfigDict(extra="forbid")

    cells: list[VisualCell] = Field(min_length=1)


class VisualWorkbook(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sheets: list[VisualSheet] = Field(min_length=1, max_length=100)


class VisualWorkbookExtractor:
    def __init__(self, model):
        self.model = model

    async def extract(self, image_urls: list[str]) -> VisualWorkbook:
        content = [{"type": "text", "text": TRANSCRIPTION_PROMPT}]
        content.extend({"type": "image_url", "image_url": {"url": url}} for url in image_urls)
        try:
            response = await self.model.ainvoke([{"role": "user", "content": content}])
        except Exception as exc:
            raise ModelServiceError(f"Workbook transcription failed: {exc}") from exc
        body = response.content
        if isinstance(body, list):
            body = "".join(part.get("text", "") for part in body if isinstance(part, dict))
        if not isinstance(body, str):
            raise ExtractionError("Workbook transcription returned non-text content")
        body = body.strip()
        if body.startswith("```json"):
            body = body[7:]
        if body.startswith("```"):
            body = body[3:]
        if body.endswith("```"):
            body = body[:-3]
        try:
            return VisualWorkbook.model_validate(json.loads(body.strip()))
        except (json.JSONDecodeError, ValidationError) as exc:
            raise ExtractionError("Workbook transcription must contain a nonempty sheets array") from exc


def build_workbook(transcription: VisualWorkbook, settings: Settings) -> bytes:
    """Write only transcribed source text; sheet names and styling add no data."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    total_cells = 0
    total_characters = 0
    total_covered = 0
    edge = Side(style="thin", color="D9E2EC")
    try:
        for sheet_number, source in enumerate(transcription.sheets, start=1):
            sheet = workbook.create_sheet(f"Table {sheet_number}")
            occupied: set[tuple[int, int]] = set()
            for item in source.cells:
                if not item.text.strip():
                    raise ExtractionError("Transcribed cells must contain visible text")
                last_row = item.row + item.row_span - 1
                last_column = item.column + item.column_span - 1
                if last_row > 1000 or last_column > 100:
                    raise ExtractionError("Transcribed cell span exceeds workbook limits")
                total_covered += item.row_span * item.column_span
                if total_covered > settings.max_workbook_cells * 10:
                    raise ExtractionError("Transcribed merged-cell area exceeds workbook limits")
                for row in range(item.row, last_row + 1):
                    for column in range(item.column, last_column + 1):
                        if (row, column) in occupied:
                            raise ExtractionError("Transcribed cells overlap")
                        occupied.add((row, column))
                total_cells += 1
                total_characters += len(item.text)
                if total_cells > settings.max_workbook_cells or total_characters > settings.max_workbook_characters:
                    raise ExtractionError("Transcribed workbook exceeds configured limits")
                cell = sheet.cell(item.row, item.column)
                cell.value = item.text
                cell.data_type = "s"  # A literal source value beginning '=' is not a formula.
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                existing_height = sheet.row_dimensions[item.row].height or 0
                sheet.row_dimensions[item.row].height = max(existing_height, 30, 15 * (item.text.count("\n") + 1))
                if item.row_span > 1 or item.column_span > 1:
                    sheet.merge_cells(start_row=item.row, start_column=item.column,
                                      end_row=last_row, end_column=last_column)
            for column in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(column)].width = 28
        stream = BytesIO()
        workbook.save(stream)
        data = stream.getvalue()
    finally:
        workbook.close()

    # Check the exported workbook, including literal strings and merged spans.
    reopened = load_workbook(BytesIO(data), read_only=False, data_only=False)
    try:
        if len(reopened.worksheets) != len(transcription.sheets):
            raise ExtractionError("Workbook export lost a source table")
        for sheet, source in zip(reopened.worksheets, transcription.sheets, strict=True):
            for item in source.cells:
                if sheet.cell(item.row, item.column).value != item.text:
                    raise ExtractionError("Workbook export changed a transcribed cell")
                if item.row_span > 1 or item.column_span > 1:
                    expected = (item.row, item.column, item.row + item.row_span - 1,
                                item.column + item.column_span - 1)
                    if not any((area.min_row, area.min_col, area.max_row, area.max_col) == expected
                               for area in sheet.merged_cells.ranges):
                        raise ExtractionError("Workbook export lost a visible merged cell")
    finally:
        reopened.close()
    return data
